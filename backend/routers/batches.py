"""Preview and queue bounded CSV/XLSX imports without storing the uploaded file."""
import csv
import io
import secrets
import zipfile
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from starlette.concurrency import run_in_threadpool
from xml.etree.ElementTree import ParseError
from pydantic import ValidationError
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from auth import get_client_or_operator
from database import get_db
from limiter import limiter
from models import AnalysisJob, BatchImport, Client, Content
from schemas import AnalyzeRequest
from services.rule_detector import mask_pii

router = APIRouter(prefix="/api/batches", tags=["batch-imports"])
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 50
MAX_SCANNED_ROWS = 5000
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024


async def _read_table(file: UploadFile):
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        raise HTTPException(status_code=422, detail="CSV 또는 XLSX 파일을 선택하세요.")
    payload = await file.read(MAX_FILE_BYTES + 1)
    if len(payload) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="파일은 2MB 이하여야 합니다.")
    return await run_in_threadpool(_parse_table, suffix, payload)


def _parse_table(suffix, payload):
    book = None
    try:
        if suffix == ".csv":
            iterator = csv.reader(io.StringIO(payload.decode("utf-8-sig"), newline=""))
        else:
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(item.file_size for item in entries) > MAX_UNCOMPRESSED_BYTES:
                    raise HTTPException(status_code=413, detail="XLSX 압축 해제 크기는 20MB 이하여야 합니다.")
            book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            if book.active is None:
                raise ValueError("Missing worksheet")
            if (book.active.max_column or 0) > MAX_COLUMNS or (book.active.max_row or 0) > MAX_SCANNED_ROWS:
                raise HTTPException(status_code=413, detail="최대 50열·5,000개 물리 행까지 읽을 수 있습니다.")
            iterator = book.active.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        if not headers or any(not header for header in headers) or len(set(headers)) != len(headers):
            raise HTTPException(status_code=422, detail="첫 행의 열 이름은 비어 있거나 중복될 수 없습니다.")
        if len(headers) > MAX_COLUMNS or any(len(header) > 100 for header in headers):
            raise HTTPException(status_code=413, detail="최대 50열이며 열 이름은 100자 이하여야 합니다.")
        rows = []
        for number, values in enumerate(iterator, 2):
            if number > MAX_SCANNED_ROWS:
                raise HTTPException(status_code=413, detail="빈 행을 포함해 최대 5,000행까지 읽을 수 있습니다.")
            if not any(value is not None and str(value).strip() for value in values):
                continue
            if len(rows) >= MAX_ROWS:
                raise HTTPException(status_code=413, detail="한 번에 최대 500행까지 접수할 수 있습니다.")
            if len(values) > len(headers) or any(len(str(value or "")) > 8000 for value in values):
                raise HTTPException(status_code=422, detail=f"{number}행의 열 개수 또는 셀 길이(최대 8,000자)를 확인하세요.")
            row = dict(zip(headers, ["" if value is None else str(value) for value in values]))
            rows.append((number, row))
        return suffix[1:], headers, rows
    except (UnicodeError, csv.Error, ValueError, StopIteration, OSError, KeyError, ParseError, zipfile.BadZipFile) as exc:
        raise HTTPException(status_code=422, detail="파일을 읽을 수 없습니다. UTF-8 CSV 또는 정상적인 XLSX를 확인하세요.") from exc
    finally:
        if book:
            book.close()


@router.post("/preview")
@limiter.limit("10/hour")
async def preview(request: Request, file: UploadFile = File(...),
                  client: Client | None = Depends(get_client_or_operator)):
    source, headers, rows = await _read_table(file)
    return {"format": source, "columns": headers, "rows": len(rows),
            "sample": [{key: mask_pii(str(value or ""))[0][:100] for key, value in row.items() if key}
                       for _, row in rows[:5]]}


@router.post("", status_code=202)
@limiter.limit("5/hour")
async def submit(request: Request, file: UploadFile = File(...),
                 content_id_column: str = Form(...), text_column: str = Form(...),
                 client: Client | None = Depends(get_client_or_operator),
                 db: Session = Depends(get_db)):
    source, headers, rows = await _read_table(file)
    if content_id_column not in headers or text_column not in headers or content_id_column == text_column:
        raise HTTPException(status_code=422, detail="서로 다른 콘텐츠 ID 열과 텍스트 열을 선택하세요.")
    client_id = client.id if client else None
    seen = set()
    candidates = []
    errors = []
    for number, row in rows:
        try:
            body = AnalyzeRequest(content_id=str(row.get(content_id_column) or "").strip(),
                                  text=str(row.get(text_column) or ""))
        except ValidationError:
            errors.append({"row": number, "reason": "콘텐츠 ID 또는 텍스트 형식이 올바르지 않습니다."})
            continue
        if body.content_id in seen:
            errors.append({"row": number, "reason": "파일 안에서 콘텐츠 ID가 중복되었습니다."})
            continue
        seen.add(body.content_id)
        masked, pii_types = mask_pii(body.text)
        candidates.append((number, body.content_id, masked, pii_types))
    ids = [item[1] for item in candidates]
    existing_contents = {value[0] for value in db.query(Content.content_id).filter(
        Content.client_id == client_id, Content.content_id.in_(ids)).all()} if ids else set()
    keys = [f"client:{client_id if client_id is not None else 'operator'}:{item}" for item in ids]
    existing_keys = {value[0] for value in db.query(AnalysisJob.submission_key).filter(
        AnalysisJob.submission_key.in_(keys)).all()} if keys else set()
    accepted = []
    for number, content_id, masked, pii_types in candidates:
        key = f"client:{client_id if client_id is not None else 'operator'}:{content_id}"
        if content_id in existing_contents or key in existing_keys:
            errors.append({"row": number, "reason": "이미 접수되었거나 분석된 콘텐츠 ID입니다."})
        else:
            accepted.append((content_id, masked, pii_types, key))
    if not rows:
        raise HTTPException(status_code=422, detail="파일에 접수할 행이 없습니다.")
    batch = BatchImport(id=secrets.token_hex(16), client_id=client_id, source_format=source,
                        rows_total=len(rows), accepted=len(accepted), skipped=len(rows) - len(accepted), errors=errors)
    db.add(batch)
    db.flush()
    for content_id, masked, pii_types, key in accepted:
        db.add(AnalysisJob(id=secrets.token_hex(16), batch_id=batch.id, client_id=client_id,
                           submission_key=key, content_id=content_id, text=masked, pii_types=pii_types,
                           kind="new", status="PENDING", attempts=0, next_attempt_at=datetime.utcnow()))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="다른 요청이 같은 콘텐츠 ID를 먼저 접수했습니다. 다시 미리보기를 실행하세요.")
    return {"id": batch.id, "rows_total": batch.rows_total, "accepted": batch.accepted,
            "skipped": batch.skipped, "errors": errors}


def _batch_result(db: Session, batch: BatchImport) -> dict:
    counts = {status: count for status, count in db.query(AnalysisJob.status, func.count())
              .filter(AnalysisJob.batch_id == batch.id).group_by(AnalysisJob.status).all()}
    return {"id": batch.id, "format": batch.source_format, "rows_total": batch.rows_total,
            "accepted": batch.accepted, "skipped": batch.skipped, "by_status": counts,
            "created_at": batch.created_at}


@router.get("")
def list_batches(response: Response, limit: int = Query(20, ge=1, le=100), offset: int = Query(0, ge=0), db: Session = Depends(get_db),
                 client: Client | None = Depends(get_client_or_operator)):
    query = db.query(BatchImport)
    if client:
        query = query.filter(BatchImport.client_id == client.id)
    response.headers["X-Total-Count"] = str(query.count())
    return [_batch_result(db, batch) for batch in query.order_by(BatchImport.created_at.desc(), BatchImport.id.desc()).offset(offset).limit(limit).all()]


@router.get("/{batch_id}")
def get_batch(batch_id: str, db: Session = Depends(get_db),
              client: Client | None = Depends(get_client_or_operator)):
    batch = db.get(BatchImport, batch_id)
    if not batch or (client and batch.client_id != client.id):
        raise HTTPException(status_code=404, detail="접수 내역이 없습니다.")
    return _batch_result(db, batch)


@router.get("/{batch_id}/errors.csv")
def download_errors(batch_id: str, db: Session = Depends(get_db),
                    client: Client | None = Depends(get_client_or_operator)):
    batch = db.get(BatchImport, batch_id)
    if not batch or (client and batch.client_id != client.id):
        raise HTTPException(status_code=404, detail="접수 내역이 없습니다.")
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["row", "reason"])
    for item in batch.errors or []:
        writer.writerow([item["row"], item["reason"]])
    return Response(stream.getvalue().encode("utf-8-sig"), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="batch-{batch.id}-errors.csv"'})
